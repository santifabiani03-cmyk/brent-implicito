"""Parámetros de mercado para el WACC y el precio argentino, todos de fuente pública:
- rf: UST 10 años (FRED DGS10).
- ERP madura, prima de riesgo país (CRP) y tasa marginal de impuesto por país: Damodaran, ctryprem.xlsx.
- Beta desapalancada sectorial: Damodaran, betas.xls ('Oil/Gas (Production and Exploration)' / '(Integrated)').
- Precio Medanito y Brent mensual (USD/m3) + regalías efectivas de Neuquén: Secretaría de Energía, Informe de Regalías."""
import csv, io, zipfile
from .common import get, guardar, ahora_iso

DAMO = "https://pages.stern.nyu.edu/~adamodar/pc/datasets/"
PAISES = {"Argentina": "Argentina", "Brasil": "Brazil", "Norteamérica": "United States"}
REG = ("http://www.energia.gob.ar/contenidos/archivos/Reorganizacion/informacion_del_mercado/mercado_hidrocarburos/"
       "informacion_estadistica/regalias/Regalias_CRUDO.zip")
M3_BBL = 6.2898108


def rf() -> dict:
    filas = list(csv.reader(io.StringIO(get("https://fred.stlouisfed.org/graph/fredgraph.csv", params={"id": "DGS10"}).text)))[1:]
    f, v = [r for r in filas if r[1] not in (".", "")][-1]
    return {"valor": float(v) / 100, "fecha": f, "fuente": "FRED DGS10 (UST 10 años)"}


def damodaran(sufijo: str = "") -> dict:
    """sufijo '' = versión vigente; '24' = archivo histórico de enero 2024 (pc/archives/)."""
    import openpyxl, xlrd
    base = DAMO if not sufijo else DAMO.replace("datasets/", "archives/")
    wb = openpyxl.load_workbook(io.BytesIO(get(base + f"ctryprem{sufijo}.xlsx", timeout=90).content), read_only=True, data_only=True)
    ws = wb["ERPs by country"]
    filas = list(ws.iter_rows(values_only=True))
    fecha = str(filas[1][1])[:10]
    erp_madura = filas[2][4]
    crp = {r[0]: r[5] for r in filas[8:] if r and isinstance(r[0], str)}
    tax = {r[0]: r[1] for r in wb["Country Tax Rates"].iter_rows(values_only=True) if r and isinstance(r[0], str) and isinstance(r[1], (int, float))}
    b = xlrd.open_workbook(file_contents=get(base + f"betas{sufijo}.xls", timeout=90).content).sheet_by_name("Industry Averages")
    betas = {}
    for i in range(b.nrows):
        nom = b.row_values(i)[0]
        if nom in ("Oil/Gas (Production and Exploration)", "Oil/Gas (Integrated)"):
            betas[nom] = b.row_values(i)[7]      # beta desapalancada corregida por caja
    v = xlrd.open_workbook(file_contents=get(base + f"vebitda{sufijo}.xls", timeout=90).content)
    vs = v.sheet_by_name("Industry Averages") if "Industry Averages" in v.sheet_names() else v.sheet_by_index(0)
    ev_ebitda = next((vs.row_values(i)[3] for i in range(vs.nrows) if vs.row_values(i)[0] == "Oil/Gas (Integrated)"), None)
    fuente = f"Damodaran (NYU Stern), ctryprem.xlsx, betas.xls y vebitda.xls, actualización {fecha}"
    return {"ev_ebitda_integradas": {"valor": ev_ebitda, "fuente": fuente + " — EV/EBITDA Oil/Gas (Integrated)"},
            "erp": {"valor": erp_madura, "fuente": fuente + " — ERP de mercado maduro"},
            "crp": {reg: {"valor": crp.get(p), "fuente": fuente + f" — CRP {p}"} for reg, p in PAISES.items()},
            "tasa_marginal": {reg: {"valor": tax.get(p), "fuente": fuente + f" — tasa marginal {p}"} for reg, p in PAISES.items()},
            "beta_u": {"pure-play": {"valor": betas.get("Oil/Gas (Production and Exploration)"), "fuente": fuente + " — Oil/Gas (Production and Exploration)"},
                       "integrada": {"valor": betas.get("Oil/Gas (Integrated)"), "fuente": fuente + " — Oil/Gas (Integrated)"}}}


def argentina() -> dict:
    import openpyxl
    z = zipfile.ZipFile(io.BytesIO(get(REG, timeout=120).content))
    wb = openpyxl.load_workbook(io.BytesIO(z.read(z.namelist()[0])), read_only=True, data_only=True)
    def tabla(hoja, cols):
        filas = list(wb[hoja].iter_rows(values_only=True))
        h = next(i for i, r in enumerate(filas) if r and r[0] == "AÑO")
        idx = {c: filas[h].index(c) for c in cols}
        out, anio = {}, None
        for r in filas[h + 1:]:
            if r[0] is not None:
                anio = r[0] if isinstance(r[0], int) else anio
            if r[1] in (None, "") or anio is None:
                continue
            try:
                out[f"{anio}-{int(r[1]):02d}"] = {c: float(r[i] or 0) for c, i in idx.items()}
            except (TypeError, ValueError):
                pass
        return out
    precios = tabla("Tabla precios", ["MEDANITO", "BRENT"])
    prod, reg = tabla("Tabla producción", ["NEUQUEN"]), tabla("Tabla regalías", ["NEUQUEN"])
    meses = sorted(m for m, v in precios.items() if v["MEDANITO"] > 0)[:-2]    # los 2 últimos son provisorios
    serie = []
    for m in meses:
        p = precios[m]
        valor_prod = prod.get(m, {}).get("NEUQUEN", 0) * p["MEDANITO"]
        serie.append({"mes": m, "medanito_usd_bbl": p["MEDANITO"] / M3_BBL, "brent_usd_bbl": p["BRENT"] / M3_BBL,
                      "regalias_pct_neuquen": (reg.get(m, {}).get("NEUQUEN", 0) / valor_prod) if valor_prod else None})
    return {"serie": serie[-24:], "fuente": REG + " (Informe de Regalías, SE; últimos 2 meses provisorios excluidos)"}


def main():
    out = {"actualizado": ahora_iso(), "errores": {}}
    for k, fn in (("rf", rf), ("damodaran", damodaran), ("argentina", argentina)):
        try:
            out[k] = fn()
        except Exception as e:  # noqa: BLE001
            out["errores"][k] = str(e)
    guardar("macro.json", out)
    d = out.get("damodaran", {})
    print("macro: rf", out.get("rf", {}).get("valor"), "| ERP", d.get("erp", {}).get("valor"),
          "| CRP", {k: round(v["valor"], 4) for k, v in d.get("crp", {}).items() if v["valor"] is not None},
          "| beta_u", {k: round(v["valor"], 3) for k, v in d.get("beta_u", {}).items() if v["valor"]},
          "| AR últ.", out.get("argentina", {}).get("serie", [{}])[-1], "| errores", out["errores"])


if __name__ == "__main__":
    main()
